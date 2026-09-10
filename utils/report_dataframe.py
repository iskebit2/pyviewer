"""
Rapor DataFrame - Word ve HTML çıktıları için optimize edilmiş
"""

import pandas as pd
from typing import Dict, Any, Optional, List, Union
from docx.shared import Pt
from docx import Document as DocxDocument


class ReportDataFrame(pd.DataFrame):
    """Özel DataFrame - Raporlama için geliştirilmiş"""
    
    _metadata = ['custom_title', 'custom_desc', 'column_units', '_table_style']
    
    def __init__(self, *args, **kwargs):
        self.custom_title = kwargs.pop("custom_title", None)
        self.custom_desc = kwargs.pop("custom_desc", None)
        self.column_units = kwargs.pop("column_units", {})
        self._table_style = kwargs.pop("table_style", 'Table Grid')
        super().__init__(*args, **kwargs)
    
    @property
    def _constructor(self):
        def _c(*args, **kwargs):
            df = ReportDataFrame(*args, **kwargs)
            df.custom_title = self.custom_title
            df.custom_desc = self.custom_desc
            df.column_units = self.column_units
            df._table_style = self._table_style
            return df
        return _c
    
    def round_floats(self, decimals: int = 2) -> 'ReportDataFrame':
        """Float sütunları yuvarlar"""
        for col in self.select_dtypes(include=['float64', 'float32']).columns:
            self[col] = self[col].round(decimals)
        return self
    
    def format_percent(self, columns: List[str], decimals: int = 1) -> 'ReportDataFrame':
        """Yüzde formatı uygular"""
        for col in columns:
            if col in self.columns:
                self[col] = self[col].apply(lambda x: f"{x*100:.{decimals}f}%")
        return self
    
    def to_markdown(self) -> str:
        """Markdown tablosuna dönüştürür"""
        lines = []
        if self.custom_title:
            lines.append(f"## {self.custom_title}")
        if self.custom_desc:
            lines.append(f"*{self.custom_desc}*")
        
        lines.append(super().to_markdown())
        return "\n".join(lines)
    
    def _repr_html_(self) -> str:
        """HTML gösterimi - Jupyter için"""
        parts = []
        
        if self.custom_title:
            parts.append(f'<h4 style="color:#2c3e50;">{self.custom_title}</h4>')
        if self.custom_desc:
            parts.append(f'<p style="font-size:0.9em;color:#7f8c8d;">{self.custom_desc}</p>')
        
        # DataFrame HTML
        html = super()._repr_html_()
        
        # Birim satırı ekle
        if self.column_units and not self.empty:
            units_row = "<tr style='font-weight:lighter;font-style:italic;font-size:0.8em;color:#95a5a6;'>"
            units_row += "<th></th>"
            for col in self.columns:
                units_row += f"<th>{self.column_units.get(col, '')}</th>"
            units_row += "</tr>"
            html = html.replace("</thead>", units_row + "</thead>")
        
        parts.append(html)
        return "".join(parts)
    
    def save_to_docx(self, document, title: Optional[str] = None, 
                     desc: Optional[str] = None, level: int = 2) -> None:
        """Word belgesine ekler"""
        title = title or self.custom_title
        desc = desc or self.custom_desc
        
        # Belge nesnesini al
        doc = document.doc if hasattr(document, 'doc') else document
        
        # Başlık ekle
        if title:
            if hasattr(document, 'add_heading_numbered'):
                document.add_heading_numbered(title, level=level)
            else:
                doc.add_heading(title, level=level)
        
        # Açıklama ekle
        if desc:
            doc.add_paragraph(desc)
        
        if self.empty:
            doc.add_paragraph("Gösterilecek veri bulunamadı.")
            return
        
        # Tablo oluştur
        row_count = len(self) + 1  # Başlık satırı
        if self.column_units:
            row_count += 1  # Birim satırı
        
        table = doc.add_table(rows=row_count, cols=len(self.columns))
        table.style = self._table_style
        
        # Başlık satırı
        for i, col_name in enumerate(self.columns):
            cell = table.rows[0].cells[i]
            run = cell.paragraphs[0].add_run(str(col_name))
            run.bold = True
            run.font.size = Pt(9)
        
        # Birim satırı
        start_row = 1
        if self.column_units:
            for i, col_name in enumerate(self.columns):
                unit = self.column_units.get(col_name, '')
                run = table.rows[1].cells[i].paragraphs[0].add_run(unit)
                run.font.size = Pt(8)
            start_row = 2
        
        # Veri satırları
        for i, (_, row) in enumerate(self.iterrows()):
            for j, value in enumerate(row):
                cell = table.rows[start_row + i].cells[j]
                text = str(value)
                # Değer çok uzunsa kısalt
                if len(text) > 50:
                    text = text[:47] + "..."
                run = cell.paragraphs[0].add_run(text)
                run.font.size = Pt(9)
        
        doc.add_paragraph()
    
    def print(self) -> None:
        """Konsol çıktısı"""
        if self.custom_title:
            print(f"\n{self.custom_title}")
            print("-" * len(self.custom_title))
        if self.custom_desc:
            print(self.custom_desc)
            print()
        
        if self.empty:
            print("Veri yok")
            return
        
        # Başlık
        print(" | ".join(self.columns))
        if self.column_units:
            units = [self.column_units.get(col, '') for col in self.columns]
            print(" | ".join(units))
        
        # Veri
        for _, row in self.iterrows():
            print(" | ".join(str(x) for x in row.values))