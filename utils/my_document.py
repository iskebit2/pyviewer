"""
Word Belge Yönetimi - Rapor Oluşturma
"""

from typing import Optional, Dict, List, Any
from docx import Document as DocxDocument
from docx.shared import Cm, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os


class myDocument:
    """Word belgesi oluşturma ve yönetme sınıfı"""
    
    def __init__(self, filename: Optional[str] = None):
        self.doc = DocxDocument(filename) if filename else DocxDocument()
        self.heading_counters: Dict[int, int] = {1: 0, 2: 0, 3: 0}
        self._section = self.doc.sections[0]
    
    def apply_visual_settings(self, 
                             top_margin: float = 2.0,
                             bottom_margin: float = 2.0,
                             left_margin: float = 2.0,
                             right_margin: float = 1.0,
                             font_name: str = 'Calibri',
                             font_size: int = 11) -> None:
        """Görsel ayarları uygular"""
        # Sayfa kenar boşlukları
        self._section.top_margin = Cm(top_margin)
        self._section.bottom_margin = Cm(bottom_margin)
        self._section.left_margin = Cm(left_margin)
        self._section.right_margin = Cm(right_margin)
        
        # Font ayarları
        style = self.doc.styles['Normal']
        style.font.name = font_name
        style.font.size = Pt(font_size)
        
        # Başlık fontu
        for level in range(1, 4):
            heading_style = self.doc.styles[f'Heading {level}']
            heading_style.font.name = font_name
    
    def add_heading_numbered(self, text: str, level: int = 1) -> Any:
        """Numaralı başlık ekler"""
        if level not in self.heading_counters:
            self.heading_counters[level] = 0
        
        self.heading_counters[level] += 1
        
        # Alt seviyeleri sıfırla
        for lvl in range(level + 1, 4):
            if lvl in self.heading_counters:
                self.heading_counters[lvl] = 0
        
        # Numara oluştur
        numbers = [str(self.heading_counters[i]) for i in range(1, level + 1)]
        number_prefix = ".".join(numbers)
        full_heading = f"{number_prefix}. {text}"
        
        return self.doc.add_heading(full_heading, level=level)
    
    def add_paragraph(self, text: str, style: Optional[str] = None) -> Any:
        """Paragraf ekler"""
        if style:
            return self.doc.add_paragraph(text, style=style)
        return self.doc.add_paragraph(text)
    
    def add_page_break(self) -> None:
        """Sayfa sonu ekler"""
        self.doc.add_page_break()
    
    def add_table(self, data: List[List[Any]], headers: List[str],
                 title: Optional[str] = None) -> None:
        """Tablo ekler"""
        if title:
            self.add_paragraph(title)
        
        if not data:
            self.add_paragraph("Veri yok")
            return
        
        table = self.doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.style = 'Table Grid'
        
        # Başlık
        for i, header in enumerate(headers):
            cell = table.rows[0].cells[i]
            run = cell.paragraphs[0].add_run(header)
            run.bold = True
        
        # Veri
        for i, row in enumerate(data):
            for j, value in enumerate(row):
                table.rows[i + 1].cells[j].text = str(value)
    
    def add_image(self, image_path: str, width_cm: float = 15.0) -> None:
        """Resim ekler"""
        if os.path.exists(image_path):
            from docx.shared import Cm
            self.doc.add_picture(image_path, width=Cm(width_cm))
        else:
            self.add_paragraph(f"Resim bulunamadı: {image_path}")
    
    def save_report(self, path: str) -> bool:
        """Raporu kaydeder"""
        try:
            # Dizin kontrolü
            os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
            self.doc.save(path)
            print(f"✅ Rapor kaydedildi: {path}")
            return True
        except Exception as e:
            print(f"❌ Kaydetme hatası: {e}")
            return False
    
    def add_section_break(self) -> None:
        """Bölüm sonu ekler"""
        self._section = self.doc.add_section()
    
    def clear(self) -> None:
        """Belgeyi temizler"""
        for element in self.doc.element.body:
            self.doc.element.body.remove(element)
        self.heading_counters = {1: 0, 2: 0, 3: 0}
        
    def preview(self, parent):
        from utils.wordPreview import WordPreview
        return WordPreview(parent, self)
        
if __name__ == "__main__":
    import tkinter as tk
    root = tk.Tk()
    doc = myDocument()

    doc.add_heading_numbered(
        "Rüzgar Yükü Analizi",
        level=1
    )

    doc.add_paragraph(
        "TS EN 1991-1-4 kapsamında rüzgar yükü hesabı."
    )

    doc.preview(root)
    root.mainloop()