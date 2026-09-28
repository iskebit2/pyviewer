"""
ReportDataFrame
---------------
Word, HTML, Markdown ve konsol çıktıları için
pandas DataFrame tabanlı raporlama sınıfı.
"""

from typing import Optional, List

import pandas as pd
from docx.shared import Pt


class ReportDataFrame(pd.DataFrame):
    """
    Raporlama için geliştirilmiş pandas DataFrame.

    Metadata:
        custom_title   : Rapor başlığı
        custom_desc    : Rapor açıklaması
        column_units   : Sütun birimleri
        column_formats : Sütun görüntüleme formatları
        _table_style   : Word tablo stili
    """

    _metadata = [
        "custom_title",
        "custom_desc",
        "column_units",
        "column_formats",
        "_table_style",
    ]

    # ------------------------------------------------------------------
    # INIT
    # ------------------------------------------------------------------

    def __init__(self, *args, **kwargs):

        self.custom_title = kwargs.pop(
            "custom_title",
            None,
        )

        self.custom_desc = kwargs.pop(
            "custom_desc",
            None,
        )

        self.column_units = kwargs.pop(
            "column_units",
            {},
        )

        self.column_formats = kwargs.pop(
            "column_formats",
            {},
        )

        self._table_style = kwargs.pop(
            "table_style",
            "Table Grid",
        )

        super().__init__(*args, **kwargs)

    # ------------------------------------------------------------------
    # PANDAS CONSTRUCTOR
    # ------------------------------------------------------------------

    @property
    def _constructor(self):

        def _c(*args, **kwargs):

            df = ReportDataFrame(
                *args,
                **kwargs,
            )

            df.custom_title = self.custom_title
            df.custom_desc = self.custom_desc

            df.column_units = (
                self.column_units.copy()
            )

            df.column_formats = (
                self.column_formats.copy()
            )

            df._table_style = self._table_style

            return df

        return _c

    # ------------------------------------------------------------------
    # FORMAT
    # ------------------------------------------------------------------

    def _format_value(
        self,
        column,
        value,
    ) -> str:
        """
        Tek bir hücreyi column_formats bilgisine göre
        metne dönüştürür.

        Örnek:
            ".3f"  -> 1.234
            ".2f"  -> 12.35
            ".1%"  -> 15.3%
            "d"    -> 3
        """

        if pd.isna(value):
            return ""

        fmt = self.column_formats.get(column)

        if fmt is None:
            return str(value)

        try:
            return format(value, fmt)

        except (TypeError, ValueError):
            return str(value)

    def _format_row(self, row) -> list[str]:
        """Bir DataFrame satırını formatlanmış metin listesine çevirir."""

        return [
            self._format_value(
                column,
                value,
            )
            for column, value in zip(
                self.columns,
                row.values,
            )
        ]

    # ------------------------------------------------------------------
    # REAL ROUNDING
    # ------------------------------------------------------------------

    def round_numeric(
        self,
        decimals: int = 2,
    ) -> "ReportDataFrame":
        """
        Numerik verileri gerçekten yuvarlar.

        Bu metod görüntüleme formatından farklıdır.
        """

        numeric_columns = self.select_dtypes(
            include="number"
        ).columns

        for column in numeric_columns:
            self[column] = self[column].round(
                decimals
            )

        return self

    # ------------------------------------------------------------------
    # MARKDOWN
    # ------------------------------------------------------------------

    def to_markdown(
        self,
        *args,
        **kwargs,
    ) -> str:
        """Formatlanmış Markdown tablosu üretir."""

        lines = []

        if self.custom_title:
            lines.append(
                f"## {self.custom_title}"
            )

        if self.custom_desc:
            lines.append(
                f"*{self.custom_desc}*"
            )

        if self.empty:
            lines.append(
                "Gösterilecek veri bulunamadı."
            )
            return "\n".join(lines)

        # Başlık
        lines.append(
            "| " + " | ".join(
                str(col)
                for col in self.columns
            ) + " |"
        )

        # Ayraç
        lines.append(
            "| " + " | ".join(
                "---"
                for _ in self.columns
            ) + " |"
        )

        # Birim
        if self.column_units:
            lines.append(
                "| " + " | ".join(
                    self.column_units.get(
                        col,
                        "",
                    )
                    for col in self.columns
                ) + " |"
            )

        # Veri
        for _, row in self.iterrows():

            values = self._format_row(row)

            lines.append(
                "| " + " | ".join(values) + " |"
            )

        return "\n".join(lines)

    # ------------------------------------------------------------------
    # HTML
    # ------------------------------------------------------------------

    def _repr_html_(self) -> str:
        """Jupyter / HTML gösterimi."""

        parts = []

        if self.custom_title:
            parts.append(
                f"<h4>{self.custom_title}</h4>"
            )

        if self.custom_desc:
            parts.append(
                f"<p>{self.custom_desc}</p>"
            )

        if self.empty:
            parts.append(
                "<p>Gösterilecek veri bulunamadı.</p>"
            )
            return "".join(parts)

        html = "<table>"

        # --------------------------------------------------------------
        # Header
        # --------------------------------------------------------------

        html += "<thead><tr>"

        for column in self.columns:
            html += (
                f"<th>{column}</th>"
            )

        html += "</tr>"

        # --------------------------------------------------------------
        # Units
        # --------------------------------------------------------------

        if self.column_units:

            html += (
                "<tr>"
            )

            for column in self.columns:

                unit = self.column_units.get(
                    column,
                    "",
                )

                html += (
                    f"<th>{unit}</th>"
                )

            html += "</tr>"

        html += "</thead>"

        # --------------------------------------------------------------
        # Data
        # --------------------------------------------------------------

        html += "<tbody>"

        for _, row in self.iterrows():

            html += "<tr>"

            for column, value in zip(
                self.columns,
                row.values,
            ):

                text = self._format_value(
                    column,
                    value,
                )

                html += (
                    f"<td>{text}</td>"
                )

            html += "</tr>"

        html += "</tbody></table>"

        parts.append(html)

        return "".join(parts)

    # ------------------------------------------------------------------
    # WORD
    # ------------------------------------------------------------------

    def save_to_docx(
        self,
        document,
        title: Optional[str] = None,
        desc: Optional[str] = None,
        level: int = 2,
    ) -> None:
        """
        Word belgesine tablo olarak ekler.
        """

        title = (
            title
            if title is not None
            else self.custom_title
        )

        desc = (
            desc
            if desc is not None
            else self.custom_desc
        )

        # --------------------------------------------------------------
        # Document
        # --------------------------------------------------------------

        doc = (
            document.doc
            if hasattr(document, "doc")
            else document
        )

        # --------------------------------------------------------------
        # Title
        # --------------------------------------------------------------

        if title:

            if hasattr(
                document,
                "add_heading_numbered",
            ):
                document.add_heading_numbered(
                    title,
                    level=level,
                )
            else:
                doc.add_heading(
                    title,
                    level=level,
                )

        # --------------------------------------------------------------
        # Description
        # --------------------------------------------------------------

        if desc:
            doc.add_paragraph(desc)

        # --------------------------------------------------------------
        # Empty
        # --------------------------------------------------------------

        if self.empty:

            doc.add_paragraph(
                "Gösterilecek veri bulunamadı."
            )

            return

        # --------------------------------------------------------------
        # Table size
        # --------------------------------------------------------------

        row_count = len(self) + 1

        if self.column_units:
            row_count += 1

        table = doc.add_table(
            rows=row_count,
            cols=len(self.columns),
        )

        table.style = self._table_style

        # --------------------------------------------------------------
        # Header
        # --------------------------------------------------------------

        for i, column in enumerate(
            self.columns
        ):

            cell = (
                table.rows[0]
                .cells[i]
            )

            run = (
                cell.paragraphs[0]
                .add_run(str(column))
            )

            run.bold = True
            run.font.size = Pt(9)

        # --------------------------------------------------------------
        # Units
        # --------------------------------------------------------------

        start_row = 1

        if self.column_units:

            for i, column in enumerate(
                self.columns
            ):

                unit = self.column_units.get(
                    column,
                    "",
                )

                run = (
                    table.rows[1]
                    .cells[i]
                    .paragraphs[0]
                    .add_run(unit)
                )

                run.font.size = Pt(8)

            start_row = 2

        # --------------------------------------------------------------
        # Data
        # --------------------------------------------------------------

        for i, (_, row) in enumerate(
            self.iterrows()
        ):

            for j, (
                column,
                value,
            ) in enumerate(
                zip(
                    self.columns,
                    row.values,
                )
            ):

                cell = (
                    table.rows[
                        start_row + i
                    ]
                    .cells[j]
                )

                text = self._format_value(
                    column,
                    value,
                )

                if len(text) > 50:
                    text = (
                        text[:47]
                        + "..."
                    )

                run = (
                    cell.paragraphs[0]
                    .add_run(text)
                )

                run.font.size = Pt(9)

        doc.add_paragraph()

    # ------------------------------------------------------------------
    # PRINT
    # ------------------------------------------------------------------

    def print(self) -> None:
        """Formatlanmış konsol çıktısı."""

        if self.custom_title:

            print(
                f"\n{self.custom_title}"
            )

            print(
                "-" * len(
                    self.custom_title
                )
            )

        if self.custom_desc:

            print(
                self.custom_desc
            )

            print()

        if self.empty:

            print(
                "Veri yok"
            )

            return

        # --------------------------------------------------------------
        # Header
        # --------------------------------------------------------------

        print(
            " | ".join(
                str(column)
                for column in self.columns
            )
        )

        # --------------------------------------------------------------
        # Units
        # --------------------------------------------------------------

        if self.column_units:

            print(
                " | ".join(
                    self.column_units.get(
                        column,
                        "",
                    )
                    for column in self.columns
                )
            )

        # --------------------------------------------------------------
        # Data
        # --------------------------------------------------------------

        for _, row in self.iterrows():

            print(
                " | ".join(
                    self._format_row(row)
                )
            )


if __name__ == "__main__":

    rows= {'Bölge': {1: 'A',
  3: 'A',
  2: 'B',
  4: 'B',
  5: 'D',
  0: 'E',
  6: 'F',
  7: 'F',
  8: 'G',
  9: 'H'},
 'Yüzey': {1: 'D2',
  3: 'D3',
  2: 'D2',
  4: 'D3',
  5: 'D4',
  0: 'D1',
  6: 'C1',
  7: 'C1',
  8: 'C1',
  9: 'C1'},
 'Tablo': {1: 'WALL',
  3: 'WALL',
  2: 'WALL',
  4: 'WALL',
  5: 'WALL',
  0: 'WALL',
  6: 'MONOPITCH',
  7: 'MONOPITCH',
  8: 'MONOPITCH',
  9: 'MONOPITCH'},
 'Yön': {1: 0, 3: 0, 2: 0, 4: 0, 5: 0, 0: 0, 6: 180, 7: 180, 8: 180, 9: 180},
 'Eğim': {1: 0.0,
  3: 0.0,
  2: 0.0,
  4: 0.0,
  5: 0.0,
  0: 0.0,
  6: 7.125016348901757,
  7: 7.125016348901757,
  8: 7.125016348901757,
  9: 7.125016348901757},
 'Cpe,10 min': {1: -1.2,
  3: -1.2,
  2: -0.8,
  4: -0.8,
  5: 0.7,
  0: -0.3,
  6: -2.38500065395607,
  7: -2.38500065395607,
  8: -1.3,
  9: -0.8425003269780351},
 'Cpe,10 max': {1: 0.0,
  3: 0.0,
  2: 0.0,
  4: 0.0,
  5: 0.0,
  0: 0.0,
  6: 0.0,
  7: 0.0,
  8: 0.0,
  9: 0.0},
 'Cpe,1 min': {1: -1.4,
  3: -1.4,
  2: -1.1,
  4: -1.1,
  5: 1.0,
  0: -0.3,
  6: -2.6275009809341054,
  7: -2.6275009809341054,
  8: -2.0,
  9: -1.2},
 'Cpe,1 max': {1: 0.0,
  3: 0.0,
  2: 0.0,
  4: 0.0,
  5: 0.0,
  0: 0.0,
  6: 0.0,
  7: 0.0,
  8: 0.0,
  9: 0.0}}
    cpe_report = ReportDataFrame(
    rows,
    custom_title="Dış Basınç Katsayıları",
    custom_desc="Rüzgar bölgelerine ait dış basınç katsayıları.",
    column_units={
        "Bölge": "",
        "Yüzey": "",
        "Tablo": "",
        "Yön": "°",
        "Eğim": "°",
        "Cpe,10 min": "",
        "Cpe,10 max": "",
        "Cpe,1 min": "",
        "Cpe,1 max": "",
    },
    column_formats={
        "Yön": ".0f",
        "Eğim": ".2f",
        "Cpe,10 min": ".3f",
        "Cpe,10 max": ".3f",
        "Cpe,1 min": ".3f",
        "Cpe,1 max": ".3f",
    },
)

    cpe_report.print()